"""Shared helpers for turning NHIS layout files into the intermediate representation (IR).

The IR for one dataset is a dict:

    {"id": ..., "source": [...], "tables": [table, ...]}
    table = {"id", "title_ko", "source", "variables": [variable, ...]}
    variable = {"name", "periods": [period, ...]}
    period = {"years", "label_ko", "type", "desc_ko": [...], "notes_ko": [...],
              "codes": [{"code", "label_ko"}], "code_ref": optional name of a shared code list}

A variable appears once per table. When its label or codes differ between
periods of the source layout, it carries several periods.
"""

from __future__ import annotations

import re

# A code token followed by a colon, at the start of a line or after a separator.
CODE_RE = re.compile(
    r"(?:(?<=^)|(?<=[\s;,/(]))"
    r"([0-9]{1,3}(?:\.[0-9])?|[A-Za-z]{1,2}[0-9]?|결측|[$~.\-])"
    r"\s*[:：]\s*"
)
VAR_RE = re.compile(r"^[A-Z][A-Z0-9_]*[A-Z0-9]$|^[A-Z]$")


def clean(text) -> str:
    if text is None:
        return ""
    return " ".join(str(text).replace("　", " ").split())


def split_codes(text: str) -> list[dict]:
    """Return codes found in one line such as '1: 없음 2: 있음' or '1:예;2:아니오;'."""
    text = clean(text)
    matches = list(CODE_RE.finditer(text))
    if not matches or matches[0].start() > 0:
        return []
    codes = []
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        label = text[m.end():end].strip(" ;,/")
        code = m.group(1)
        if code == "결측":
            code = "(missing)"
        codes.append({"code": code, "label_ko": label})
    return codes


def classify_lines(lines: list[str]) -> tuple[list[str], list[str], list[dict]]:
    """Split value-description lines into descriptions, notes, and codes."""
    desc, notes, codes = [], [], []
    for raw in lines:
        line = clean(raw)
        if not line or line in ("-", "〃"):
            continue
        if line.startswith(("※", "*")):
            notes.append(line.lstrip("※* ").strip())
            continue
        found = split_codes(line)
        if found:
            codes.extend(found)
        else:
            desc.append(line)
    # Drop duplicate codes that come from merged cells repeated across rows.
    seen, unique = set(), []
    for c in codes:
        key = (c["code"], c["label_ko"])
        if key not in seen:
            seen.add(key)
            unique.append(c)
    return dedupe(desc), dedupe(notes), unique


def dedupe(items: list[str]) -> list[str]:
    out = []
    for x in items:
        if x not in out:
            out.append(x)
    return out


def new_period(years: str, label_ko: str, type_: str = "", lines=None, codes=None) -> dict:
    desc, notes, found = classify_lines(lines or [])
    return {
        "years": years,
        "label_ko": clean(label_ko),
        "type": type_,
        "desc_ko": desc,
        "notes_ko": notes,
        "codes": (codes or []) + found,
    }


def add_variable(table: dict, name: str, period: dict) -> None:
    """Add a period to a variable, merging identical periods across years."""
    for v in table["variables"]:
        if v["name"] == name:
            for p in v["periods"]:
                if same_content(p, period):
                    p["years"] = merge_years(p["years"], period["years"])
                    return
            v["periods"].append(period)
            return
    table["variables"].append({"name": name, "periods": [period]})


def same_content(a: dict, b: dict) -> bool:
    keys = ("label_ko", "codes", "desc_ko", "notes_ko", "code_ref")
    return all(a.get(k) == b.get(k) for k in keys)


def merge_years(a: str, b: str) -> str:
    """Merge two year ranges such as '2002-2004' and '2005' into '2002-2005' when contiguous."""
    def bounds(s):
        parts = re.findall(r"\d{4}", s)
        return (int(parts[0]), int(parts[-1])) if parts else None

    ra, rb = bounds(a), bounds(b)
    if not ra or not rb:
        return a if a == b else f"{a}, {b}"
    lo, hi = sorted([ra, rb])
    if hi[0] <= lo[1] + 1:
        start, end = lo[0], max(lo[1], hi[1])
        return f"{start}" if start == end else f"{start}-{end}"
    return f"{a}, {b}"


def read_sheet(path, sheet, fill_merged: bool = True) -> list[list[str]]:
    """Read a worksheet into rows of cleaned strings, filling merged cells if asked."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet]
    merged = {}
    if fill_merged:
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        merged[(r, c)] = v
    rows = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None:
                v = merged.get((r, c))
            row.append(clean(v))
        rows.append(row)
    return rows


def sheet_names(path) -> list[str]:
    import openpyxl

    return openpyxl.load_workbook(path, read_only=True).sheetnames


def norm_years(text: str) -> str:
    text = clean(text).replace("~", "-").replace("–", "-").replace("년", "")
    parts = re.findall(r"\d{4}", text)
    if len(parts) >= 2:
        return f"{parts[0]}-{parts[-1]}"
    if parts:
        return parts[0]
    return text
